"""Table exports (spec §7 `/exports`, and §6: scoping must cover exports, not just lists).

Why XLSX and not PDF yet: the sheet is what the operations team actually re-works (filters, pivots),
and openpyxl is already a dependency for the insured import. A PDF of Arabic text additionally needs
shaping + an embedded font, which is a separate decision — see docs/decisions/010-exports.md.

Everything here reads through the same scope predicates as the list endpoints (`with_scope`), so a
report export can never contain rows the caller could not see on screen.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from enum import StrEnum
from io import BytesIO
from typing import Any
from uuid import UUID

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, String, and_, case, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.scope import Scope
from app.domain.enums import RequestStatus
from app.models.auth import User
from app.models.clients import Client, InsuredPerson
from app.models.masters import InsuranceType, Insurer
from app.models.requests import ServiceRequest
from app.repositories.scope_filters import with_scope
from app.services import reports

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_OPEN = [RequestStatus.NEW, RequestStatus.IN_PROGRESS, RequestStatus.WAITING_INSURER]
EXPORT_CAP = 5000
_HEADER_FILL = PatternFill("solid", fgColor="EEF0FF")
_DATE_FORMAT = "yyyy/mm/dd"
_DATETIME_FORMAT = "yyyy/mm/dd hh:mm"


class ExportKind(StrEnum):
    REQUESTS = "requests"
    CLIENTS = "clients"
    INSURER_PERFORMANCE = "insurer-performance"


@dataclass(frozen=True)
class Column:
    header: str
    key: str
    width: int = 18
    number_format: str | None = None


@dataclass(frozen=True)
class ExportResult:
    filename: str
    media_type: str
    content: bytes
    rows: int
    truncated: bool
    sheet: str


COLUMNS: dict[ExportKind, list[Column]] = {
    ExportKind.REQUESTS: [
        Column("رقم الطلب", "number", 12),
        Column("الحالة", "status", 14),
        Column("الأولوية", "priority", 12),
        Column("القناة", "channel", 12),
        Column("الموظف", "employee", 22),
        Column("الشركة العميلة", "client", 26),
        Column("نوع التأمين", "insurance_type", 14),
        Column("شركة التأمين", "insurer", 26),
        Column("المسؤول", "assignee", 20),
        Column("متأخر", "is_overdue", 10),
        Column("أول رد بعد (دقيقة)", "first_response_minutes", 18),
        Column("فتح في", "created_at", 18, _DATETIME_FORMAT),
        Column("استحقاق SLA", "sla_due_at", 18, _DATETIME_FORMAT),
    ],
    ExportKind.CLIENTS: [
        Column("الشركة", "name", 30),
        Column("النوع", "type", 12),
        Column("الحالة", "status", 14),
        Column("الموظفون المؤمَّن عليهم", "insured_count", 20),
        Column("طلبات مفتوحة", "open_requests", 14),
        Column("المسؤول", "account_owner", 22),
        Column("عميل منذ", "since_date", 16, _DATE_FORMAT),
        Column("المصدر", "source", 22),
    ],
    ExportKind.INSURER_PERFORMANCE: [
        Column("شركة التأمين", "insurer", 32),
        Column("الطلبات", "requests", 12),
        Column("متوسط أول رد (دقيقة)", "avg_first_response_minutes", 20),
        Column("الالتزام بالـSLA %", "sla_pct", 18),
        Column("التقييم", "tone", 14),
    ],
}

SHEET_TITLE = {
    ExportKind.REQUESTS: "الطلبات",
    ExportKind.CLIENTS: "العملاء",
    ExportKind.INSURER_PERFORMANCE: "أداء الشركات",
}

_TONE_LABEL = {"g": "ممتاز", "a": "مقبول", "r": "يحتاج تدخلاً"}


def _requests_statement(
    *,
    statuses: list[RequestStatus] | None = None,
    only_open: bool = False,
    only_overdue: bool = False,
    q: str | None = None,
    client_id: UUID | None = None,
) -> Select[tuple[Any, ...]]:
    # PostgreSQL has no round(double precision, integer): the epoch arithmetic must be cast to numeric
    # first, which is also what keeps the SQLite (test) and PostgreSQL (production) answers identical.
    elapsed_minutes = sa.cast(
        (func.extract("epoch", ServiceRequest.first_response_at) - func.extract("epoch", ServiceRequest.created_at)) / 60.0,
        sa.Numeric(12, 4),
    )
    response_minutes = func.round(elapsed_minutes, 1)
    stmt = (
        select(
            ServiceRequest.number.label("number"),
            ServiceRequest.status.label("status"),
            ServiceRequest.priority.label("priority"),
            ServiceRequest.channel.label("channel"),
            InsuredPerson.name.label("employee"),
            Client.name.label("client"),
            InsuranceType.code.label("insurance_type"),
            Insurer.name.label("insurer"),
            User.full_name.label("assignee"),
            ServiceRequest.is_overdue.label("is_overdue"),
            response_minutes.label("first_response_minutes"),
            ServiceRequest.created_at.label("created_at"),
            ServiceRequest.sla_due_at.label("sla_due_at"),
        )
        .outerjoin(InsuredPerson, InsuredPerson.id == ServiceRequest.insured_person_id)
        .outerjoin(Client, Client.id == ServiceRequest.client_id)
        .outerjoin(InsuranceType, InsuranceType.id == ServiceRequest.insurance_type_id)
        .outerjoin(Insurer, Insurer.id == ServiceRequest.insurer_id)
        .outerjoin(User, User.id == ServiceRequest.assigned_to)
    )
    conds: list[Any] = []
    if statuses:
        conds.append(ServiceRequest.status.in_(statuses))
    if only_open:
        conds.append(ServiceRequest.status.in_(_OPEN))
    if only_overdue:
        conds.append(ServiceRequest.is_overdue.is_(True))
    if client_id is not None:
        conds.append(ServiceRequest.client_id == client_id)
    if q:
        like = f"%{q.strip().lower()}%"
        conds.append(
            or_(
                func.lower(Client.name).like(like),
                func.lower(InsuredPerson.name).like(like),
                func.lower(ServiceRequest.subject).like(like),
                func.lower(func.cast(ServiceRequest.number, String(20))).like(like),
            )
        )
    if conds:
        stmt = stmt.where(and_(*conds))
    return stmt.order_by(ServiceRequest.number.desc())


def _clients_statement() -> Select[tuple[Any, ...]]:
    insured_count = (
        select(func.count(InsuredPerson.id))
        .where(InsuredPerson.client_id == Client.id, InsuredPerson.deleted_at.is_(None))
        .scalar_subquery()
    )
    open_requests = (
        select(func.count(ServiceRequest.id))
        .where(
            ServiceRequest.client_id == Client.id,
            ServiceRequest.status.in_([RequestStatus.NEW, RequestStatus.IN_PROGRESS, RequestStatus.WAITING_INSURER]),
        )
        .scalar_subquery()
    )
    return select(
        Client.name.label("name"),
        Client.type.label("type"),
        Client.status.label("status"),
        insured_count.label("insured_count"),
        open_requests.label("open_requests"),
        User.full_name.label("account_owner"),
        Client.since_date.label("since_date"),
        case((Client.source_lead_id.is_not(None), "lead"), else_="manual").label("source"),
    ).outerjoin(User, User.id == Client.account_owner_id)


async def _rows(
    kind: ExportKind,
    session: AsyncSession,
    scope: Scope,
    *,
    statuses: list[RequestStatus] | None = None,
    only_open: bool = False,
    only_overdue: bool = False,
    q: str | None = None,
    client_id: UUID | None = None,
) -> list[dict[str, Any]]:
    if kind is ExportKind.INSURER_PERFORMANCE:
        return await reports.insurer_performance(session, scope)

    if kind is ExportKind.REQUESTS:
        stmt = with_scope(
            _requests_statement(statuses=statuses, only_open=only_open, only_overdue=only_overdue, q=q, client_id=client_id),
            scope,
            "request",
        )
    else:
        stmt = with_scope(_clients_statement(), scope, "client")
    rows = (await session.execute(stmt.limit(EXPORT_CAP + 1))).mappings().all()
    out: list[dict[str, Any]] = []
    for row in rows:
        item = dict(row)
        for key, value in list(item.items()):
            if hasattr(value, "value"):  # enum columns: export the wire value, not the Python name
                item[key] = value.value
        if kind is ExportKind.CLIENTS:
            item["source"] = "تحول من Lead" if item.get("source") == "lead" else "مسجّل يدوياً"
        out.append(item)
    return out


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return "نعم" if value else "لا"
    if isinstance(value, datetime | date):
        return value
    if isinstance(value, int | float):
        return value
    return str(value)


def _workbook(sheet: str, columns: list[Column], rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.sheet_view.rightToLeft = True  # Arabic-first: the sheet opens the way the screen reads (spec §11)
    ws.freeze_panes = "A2"
    ws.append([column.header for column in columns])
    for index, cell in enumerate(ws[1], start=1):
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = columns[index - 1].width
    for row in rows:
        ws.append([_cell(row.get(column.key)) for column in columns])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            fmt = columns[cell.column - 1].number_format
            if fmt and isinstance(cell.value, datetime | date):
                cell.number_format = fmt
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def period_workbook(data: dict[str, Any]) -> bytes:
    """The ERP period report (ADR-033/E2): totals + per-type in/out/net, from the same aggregate the
    screen reads — numbers are computed by the server once, never re-added by the frontend."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Period"
    ws.sheet_view.rightToLeft = True
    ws.append(["ERP period report — ERP تشغيلي للمكتب (ليس دفاتر ضريبية رسمية)"])
    ws.append(["period", data.get("label"), str(data.get("start")), str(data.get("end"))])
    ws.append(["total_in", data.get("total_in")])
    ws.append(["total_out", data.get("total_out")])
    ws.append(["net", data.get("net")])
    ws.append([])
    ws.append(["entry_type", "in_total", "out_total", "net"])
    for row in data.get("by_type", []):
        ws.append([row.get("entry_type"), row.get("in_total"), row.get("out_total"), row.get("net")])
    for cell in ws[7]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


async def build(
    kind: ExportKind,
    session: AsyncSession,
    scope: Scope,
    *,
    statuses: list[RequestStatus] | None = None,
    only_open: bool = False,
    only_overdue: bool = False,
    q: str | None = None,
    client_id: UUID | None = None,
) -> ExportResult:
    if kind is ExportKind.INSURER_PERFORMANCE:
        # The report is an aggregation: it has no row filters, and honouring none beats pretending to.
        statuses, only_open, only_overdue, q, client_id = None, False, False, None, None
    raw = await _rows(kind, session, scope, statuses=statuses, only_open=only_open, only_overdue=only_overdue, q=q, client_id=client_id)
    truncated = len(raw) > EXPORT_CAP
    rows = raw[:EXPORT_CAP]
    columns = COLUMNS[kind]
    if kind is ExportKind.INSURER_PERFORMANCE:
        for row in rows:
            row["tone"] = _TONE_LABEL.get(str(row.get("tone")), "—")
    content = _workbook(SHEET_TITLE[kind], columns, rows)
    return ExportResult(
        filename=f"broker-hub-{kind.value}.xlsx",
        media_type=XLSX_MEDIA_TYPE,
        content=content,
        rows=len(rows),
        truncated=truncated,
        sheet=SHEET_TITLE[kind],
    )
