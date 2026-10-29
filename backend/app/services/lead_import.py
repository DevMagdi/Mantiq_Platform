"""Lead pool import from a spreadsheet: parse → dry-run preview → commit → report.

Two rules shape this module, and both come from what a manager actually does with a broker's Excel file:

* **The preview is the promise.** ``dry_run`` decides every field, including the duplicate verdict, and
  stores the resolved rows in ``rows_payload``; ``commit`` inserts exactly those and nothing else. So the
  table the manager read before pressing «تنفيذ» is the table that becomes leads.
* **The duplicate rule is not reimplemented here.** `app.services.leads.create` owns it (`find_active_by_mobile`
  + the hidden-owner case), so commit calls that service per row and turns its `Conflict` into a skip. A
  second copy of the rule in the importer would drift the day the form changes, and an import that quietly
  creates a duplicate lead is worse than one that refuses the row.

`CSV` is accepted next to `.xlsx` because that is what most brokers export; the parser is strict in the same
direction `services/clients.parse_insured_workbook` is — anything unreadable is a 422 with the reason, never
a 500.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import UTC, date, datetime
from typing import Any
from uuid import UUID

import sqlalchemy as sa
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import Conflict, NotFound, ValidationFailed
from app.core.scope import Scope
from app.domain.enums import ImportJobStatus
from app.domain.lead_fsm import TERMINAL
from app.domain.phone import normalize_mobile
from app.models.auth import Role, User
from app.models.leads import Lead, LeadImportJob
from app.models.masters import InsuranceType
from app.schemas.leads import LeadCreate
from app.services import audit
from app.services import leads as lead_svc

#: a sheet with more rows than this is a data warehouse, not an import; the cap keeps one request
#: (and one transaction) inside the API's own timeout instead of half-applying a file
MAX_ROWS = 5_000
#: the same ceiling the insured-persons import enforces (see `app/api/v1/endpoints/clients.py`)
MAX_FILE_BYTES = 5_000_000
#: what `source` says when the file does not, so the pipeline can always answer "where did this come from"
DEFAULT_SOURCE = "استيراد Excel"
#: how many resolved rows the preview carries back for the UI table
PREVIEW_ROWS = 25

#: logical column → accepted header spellings, Arabic and English. Comparison is after `_norm_header`,
#: so case, spacing, a trailing colon and the BOM do not matter.
COLUMNS: dict[str, tuple[str, ...]] = {
    "name": ("name", "full name", "الاسم", "اسم", "اسم العميل", "العميل"),
    "mobile": (
        "mobile",
        "phone",
        "mobile number",
        "phone number",
        "الموبايل",
        "موبايل",
        "التليفون",
        "رقم التليفون",
        "رقم الموبايل",
        "هاتف",
        "الجوال",
    ),
    "company": ("company", "company name", "الشركة", "شركة", "جهة العمل"),
    "position": ("position", "job title", "title", "البوزيشن", "المسمى الوظيفي", "المسمى"),
    "email": ("email", "e-mail", "البريد", "البريد الإلكتروني", "ايميل"),
    "source": ("source", "المصدر", "مصدر", "من أين"),
    "insurance_type": ("insurance type", "نوع التأمين", "نوع الوثيقة", "التأمين المطلوب"),
    "headcount": ("headcount", "employees", "number of employees", "cases", "عدد الأفراد", "عدد الموظفين", "العدد"),
    "followup": ("follow up", "follow-up", "followup", "next follow up", "تاريخ المتابعة", "المتابعة", "موعد المتابعة"),
    "rep": ("rep", "sales rep", "owner", "assigned to", "البياع", "البائع", "البياع المسؤول", "المسؤول", "المسند إلى"),
}

_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")
_DATE_PATTERNS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y")
_EMAIL_SHAPE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$")


#: Arabic headers differ by the definite article and by how the hamza was typed: «الإيميل» / «ايميل» /
#: «الايميل» are the same column to a human, and an Excel file written by a broker uses all three. Folding
#: them here means the alias list does not have to spell every variant, and a column can never be dropped
#: just because it arrived with «ال» in front (R3: a bad-email row was imported as *ready* because
#: «الإيميل» matched nothing and the whole column vanished).
_AR_FOLD = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه", "ۀ": "ه", "ـ": ""})


def _norm_header(value: Any) -> str:
    """`  الشركة : ` → `شركه`, `Mobile:` → `mobile`. Comparison is on this form only."""
    text = str(value or "").replace("\ufeff", "").strip().rstrip(":").strip().lower()
    text = re.sub(r"\s+", " ", text).translate(_AR_FOLD)
    return re.sub(r"^ال(?=\S)", "", text)


def _header_aliases() -> dict[str, str]:
    """normalised spelling → logical column, built from `COLUMNS` once."""
    out: dict[str, str] = {}
    for key, aliases in COLUMNS.items():
        for alias in aliases:
            out.setdefault(_norm_header(alias), key)
    return out


def _norm_name(value: Any) -> str:
    """Names are matched on case, spacing and digits — Arabic-Indic included, because Excel keeps them."""
    return re.sub(r"\s+", " ", str(value or "").translate(_ARABIC_DIGITS).strip()).casefold()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _phone(value: Any) -> str | None:
    """A mobile cell may hold a float, Arabic digits or a spacer's `+20 100 … `; normalize_mobile owns
    the E.164 rules, this only keeps `1.0070e+10` from becoming a number nobody can dial."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return normalize_mobile(f"{value:.0f}")
    return normalize_mobile(str(value).translate(_ARABIC_DIGITS))


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    text = text.translate(_ARABIC_DIGITS).split("T")[0].split(" ")[0]
    for pattern in _DATE_PATTERNS:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _as_int(value: Any) -> int | None:
    """`"١٢٠"` and `120.0` and `"120 employee"` all mean 120; anything else means no value."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"\d+", str(value).translate(_ARABIC_DIGITS))
    return int(match.group()) if match else None


def _followup_at(value: Any) -> tuple[datetime | None, bool]:
    """Returns (moment, unparsable). A date-only cell means *that day*, at 00:00 UTC — not "today", and
    never a silently dropped deadline: a row whose date cannot be read is reported as an error instead."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, False
    if isinstance(value, datetime):
        return (value if value.tzinfo else value.replace(tzinfo=UTC)), False
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC), False
    text = _text(value)
    if text is None:
        return None, False
    for pattern in (*_DATE_PATTERNS, "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M"):
        try:
            parsed = datetime.strptime(text.translate(_ARABIC_DIGITS), pattern)
        except ValueError:
            continue
        return (parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)), False
    return None, True


def parse_workbook(content: bytes, filename: str) -> list[dict[str, Any]]:
    """Rows as `{logical_key: cell}`, tagged with `_row` (the 1-based sheet line). Raises 422 kinds."""
    if not content:
        raise ValidationFailed("The file is empty", details={"filename": filename})
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _read_xlsx(content, filename)
    elif name.endswith((".csv", ".txt")):
        rows = _read_csv(content, filename)
    else:
        raise ValidationFailed(
            "ارفع ملف Excel بصيغة ‎.xlsx أو CSV",
            details={"filename": filename, "accepted": [".xlsx", ".csv"]},
        )
    if not rows:
        raise ValidationFailed("لا توجد صفوف تحت عناوين الأعمدة", details={"filename": filename})
    if len(rows) > MAX_ROWS:
        raise ValidationFailed(
            f"الملف أكبر من {MAX_ROWS} صفاً",
            details={"filename": filename, "rows": len(rows), "limit": MAX_ROWS},
        )
    return rows


def _read_xlsx(content: bytes, filename: str) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - the dependency is pinned in requirements
        raise ValidationFailed("Excel support is not installed") from exc
    if not content[:2] == b"PK":
        raise ValidationFailed("الملف ليس مصنف Excel بصيغة ‎.xlsx", details={"filename": filename, "hint": "CSV مقبول أيضًا"})
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a family of zip/xml errors depending on the damage
        raise ValidationFailed("تعذّر قراءة مصنف Excel", details={"filename": filename, "reason": type(exc).__name__}) from exc
    sheet = workbook.active
    return _map_header([list(r) for r in sheet.iter_rows(values_only=True)])


def _read_csv(content: bytes, filename: str) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel  # a one-column or oddly-quoted sheet still reads as comma
    return _map_header([list(row) for row in csv.reader(io.StringIO(text), dialect)])


def _map_header(grid: list[list[Any]]) -> list[dict[str, Any]]:
    """First non-empty line is the header; a missing `name` is fatal, unknown columns are reported.

    A column nobody recognises is not thrown away silently any more: its header is carried back on the first
    row (`_ignored`) so the preview can say «هذه الأعمدة لم تُقرأ» — a manager who typed an e-mail column
    should not have to discover that every address was dropped by reading the database.
    """
    lines = [row for row in grid if any(cell is not None and str(cell).strip() != "" for cell in row)]
    if not lines:
        return []
    header = [_norm_header(cell) for cell in lines[0]]
    aliases = _header_aliases()
    mapping: dict[str, int] = {}
    ignored: list[str] = []
    for idx, cell in enumerate(header):
        key = aliases.get(cell)
        if key is None:
            if cell:
                ignored.append(str(lines[0][idx]).strip())
            continue
        # first spelling wins, so two columns that mean the same thing do not overwrite each other
        mapping.setdefault(key, idx)
    if "name" not in mapping:
        raise ValidationFailed(
            "عمود «الاسم / Name» مطلوب في الصف الأول", details={"header": header, "mapped": sorted(mapping.values()) and sorted(mapping)}
        )
    rows: list[dict[str, Any]] = []
    for line_no, raw in enumerate(lines[1:], start=2):
        item: dict[str, Any] = {"_row": line_no}
        for key, idx in mapping.items():
            item[key] = raw[idx] if idx < len(raw) else None
        rows.append(item)
    if rows and ignored:
        rows[0]["_ignored"] = ignored
    return rows


async def _type_lookup(session: AsyncSession) -> dict[str, tuple[UUID, str]]:
    """Normalised name/code → (id, display name), for the "نوع التأمين" column."""
    out: dict[str, tuple[UUID, str]] = {}
    rows = (await session.execute(select(InsuranceType.id, InsuranceType.code, InsuranceType.name, InsuranceType.name_en))).all()
    for row_id, code, name, name_en in rows:
        label = str(name)
        for alias in (code, name, name_en):
            key = _norm_name(alias)
            if key:
                out[key] = (row_id, label)
    return out


async def _rep_lookup(session: AsyncSession) -> tuple[dict[str, UUID], dict[str, list[UUID]]]:
    """Active `sales_rep` users, by e-mail and by name — the same population `/sales/team` lists.

    Ambiguous names are kept as a list so a collision reports itself instead of silently picking someone.
    """
    rows = (
        await session.execute(
            select(User.id, User.full_name, User.email)
            .join(Role, Role.id == User.role_id)
            .where(func.lower(Role.code) == "sales_rep", User.is_active.is_(True), User.deleted_at.is_(None))
        )
    ).all()
    by_email: dict[str, UUID] = {}
    by_name: dict[str, list[UUID]] = {}
    for row_id, full_name, email in rows:
        by_email[str(email or "").strip().lower()] = row_id
        by_name.setdefault(_norm_name(full_name), []).append(row_id)
    return by_email, by_name


def _blank_row(row: dict[str, Any]) -> dict[str, Any]:
    return {"_row": row.get("_row"), "name": None}


async def dry_run(session: AsyncSession, actor: User, *, rows: list[dict[str, Any]], filename: str | None) -> LeadImportJob:
    """Resolve every row once, so the manager approves the real outcome instead of a hope."""
    types = await _type_lookup(session)
    rep_by_email, rep_by_name = await _rep_lookup(session)
    taken = set(
        (
            await session.execute(
                select(Lead.mobile).where(
                    Lead.deleted_at.is_(None), Lead.mobile.is_not(None), Lead.status.notin_([s.value for s in TERMINAL])
                )
            )
        ).scalars()
    )

    prepared: list[dict[str, Any]] = []
    report: list[dict[str, Any]] = []
    seen: set[str] = set()
    skipped_duplicates = 0

    ignored = (rows[0].get("_ignored") if rows else None) or []
    if ignored:
        report.append(
            {
                "row": 1,
                "field": "header",
                "message": f"أعمدة لم يتم التعرف عليها وتُركت: {'، '.join(ignored)}",
                "kind": "warning",
                "value": ignored,
            }
        )

    for row in rows:
        row_no = row.get("_row")
        payload = _blank_row(row)
        warnings: list[str] = []

        person = _text(row.get("name"))
        company = _text(row.get("company"))
        if not person and not company:
            report.append({"row": row_no, "field": "name", "message": "لا اسم ولا شركة في الصف", "kind": "error"})
            continue
        if person and len(person) > 200 and not company:
            report.append({"row": row_no, "field": "name", "message": "الاسم أطول من 200 حرف", "kind": "error"})
            continue

        # a company row keeps the account name in `name` (that is what the pipeline lists) and moves the
        # person to the contact pair; an individual row is the person
        lead_name = (company or person or "")[:200]
        payload["name"] = lead_name
        if company and person and person != company:
            payload["contact_name"] = person[:160]
            payload["type"] = "company"
        elif company:
            payload["type"] = "company"
        else:
            payload["type"] = "individual"
        title = _text(row.get("position"))
        if title:
            payload["contact_title"] = title[:120]

        raw_mobile = row.get("mobile")
        if raw_mobile is None or str(raw_mobile).strip() == "":
            report.append({"row": row_no, "field": "mobile", "message": "الموبايل مطلوب لتمييز العميل", "kind": "error"})
            continue
        mobile = _phone(raw_mobile)
        if mobile is None:
            report.append(
                {"row": row_no, "field": "mobile", "message": "رقم لا يصلح كرقم محمول مصري", "kind": "error", "value": str(raw_mobile)}
            )
            continue
        payload["mobile"] = mobile
        if mobile in seen:
            report.append({"row": row_no, "field": "mobile", "message": "مكرر داخل نفس الملف — تم تخطيه", "kind": "skipped"})
            skipped_duplicates += 1
            continue
        if mobile in taken:
            report.append({"row": row_no, "field": "mobile", "message": "يوجد lead نشط بنفس الرقم — تم تخطيه", "kind": "skipped"})
            skipped_duplicates += 1
            seen.add(mobile)
            continue
        seen.add(mobile)

        email = _text(row.get("email"))
        if email:
            if not _EMAIL_SHAPE.match(email):
                report.append({"row": row_no, "field": "email", "message": "بريد إلكتروني غير صالح", "kind": "error", "value": email})
                continue
            payload["email"] = email

        source = _text(row.get("source")) or DEFAULT_SOURCE
        payload["source"] = source[:80]
        if len(source) > 80:
            warnings.append("المصدر اقتُطع عند 80 حرفًا")

        headcount = _as_int(row.get("headcount"))
        if row.get("headcount") not in (None, "") and headcount is None:
            warnings.append("عدد الأفراد لم يُفهم — تُرك فارغًا")
        elif headcount is not None:
            payload["headcount"] = max(headcount, 1)

        moment, broken = _followup_at(row.get("followup"))
        if broken:
            report.append(
                {
                    "row": row_no,
                    "field": "followup",
                    "message": "تاريخ متابعة غير مفهوم",
                    "kind": "error",
                    "value": str(row.get("followup")),
                }
            )
            continue
        if moment is not None:
            payload["next_followup_at"] = moment.isoformat()

        type_key = _norm_name(row.get("insurance_type"))
        if type_key:
            found = types.get(type_key)
            if found is None:
                warnings.append(f"نوع التأمين «{row.get('insurance_type')}» لا يطابق أي نوع معروف — تُرك فارغًا")
            else:
                payload["insurance_type_id"] = str(found[0])
                # the preview table shows a name, not a uuid; the label is display-only, so it stays internal
                payload["_insurance_label"] = found[1]

        rep_label = _text(row.get("rep"))
        if rep_label:
            target: UUID | None = rep_by_email.get(str(rep_label).strip().lower())
            if target is None:
                matches = rep_by_name.get(_norm_name(rep_label), [])
                if len(matches) == 1:
                    target = matches[0]
                elif len(matches) > 1:
                    warnings.append(f"اسم «{rep_label}» يطابق أكثر من بياع — سيبقى غير مسند")
                else:
                    warnings.append(f"لا بياع باسم «{rep_label}» — سيبقى غير مسند")
            if target is not None:
                payload["_assign_to"] = str(target)

        if warnings:
            for note in warnings:
                report.append({"row": row_no, "field": "note", "message": note, "kind": "warning"})
        payload["_warnings"] = warnings
        prepared.append(payload)

    ready_rows = [row for row in prepared if row.get("name")]
    errors = [entry for entry in report if entry["kind"] == "error"]
    job = LeadImportJob(
        status=ImportJobStatus.DRY_RUN,
        filename=(filename or "import")[:255],
        total_rows=len(rows),
        stats={
            "rows": len(rows),
            "ready": len(ready_rows),
            "errors": len(errors),
            "skipped_duplicate": skipped_duplicates,
            "warnings": sum(1 for entry in report if entry["kind"] == "warning"),
            "assigned": sum(1 for row in ready_rows if row.get("_assign_to")),
            "unassigned": sum(1 for row in ready_rows if not row.get("_assign_to")),
        },
        error_report=report,
        rows_payload=ready_rows,
        created_by=actor.id,
    )
    session.add(job)
    await session.flush()
    return job


def brief(job: LeadImportJob) -> dict[str, Any]:
    """The wire shape for both the preview and the final report. `_`-prefixed keys stay internal."""
    preview = []
    for row in job.rows_payload[:PREVIEW_ROWS]:
        shown = {key: value for key, value in row.items() if not key.startswith("_")}
        label = row.get("_insurance_label")
        if label:
            shown["insurance_type"] = label
        preview.append(shown)
    return {
        "id": job.id,
        "status": job.status.value,
        "filename": job.filename,
        "total_rows": job.total_rows,
        "stats": job.stats,
        "error_report": job.error_report[:200],
        "preview": preview,
        "created_at": job.created_at,
        "committed_at": job.committed_at,
    }


async def get_job(session: AsyncSession, job_id: UUID) -> LeadImportJob:
    job = await session.get(LeadImportJob, job_id)
    if job is None:
        raise NotFound("Import job not found")
    return job


async def commit(session: AsyncSession, actor: User, scope: Scope, job_id: UUID, *, ip: str | None = None) -> LeadImportJob:
    """Apply a previewed job in one transaction. Re-running it is a 409, never a second set of leads."""
    job = await get_job(session, job_id)
    if job.status is not ImportJobStatus.DRY_RUN:
        raise Conflict(
            f"المهمة لم تعد معاينةً قابلة للتنفيذ — الحالة «{job.status.value}»، ولا تُقرأ صفوف الملف مرتين",
            details={"status": job.status.value, "job_id": str(job.id)},
        )
    if scope.can("leads.assign"):
        assigner_id: UUID | None = actor.id
    else:  # a viewer may not hand leads to a colleague through the side door of an import
        assigner_id = None

    job.status = ImportJobStatus.RUNNING
    await session.flush()

    created = 0
    still_duplicate = 0
    failures: list[dict[str, Any]] = []
    for row in job.rows_payload:
        payload = {key: value for key, value in row.items() if not key.startswith("_")}
        moment = payload.get("next_followup_at")
        if isinstance(moment, str):
            payload["next_followup_at"] = datetime.fromisoformat(moment)
        assign_to = row.get("_assign_to")
        if assigner_id is not None and assign_to:
            payload["assigned_to"] = UUID(str(assign_to))
        # the duplicate rule is the form's rule, not a copy of it: if a lead appeared while the manager
        # was reading the preview, this row is skipped with the same reason the form would give
        try:
            lead = await lead_svc.create(session, actor, scope, LeadCreate.model_validate(payload), ip=ip)
        except Conflict as exc:
            still_duplicate += 1
            failures.append({"row": row.get("_row"), "field": "mobile", "message": exc.message, "kind": "skipped"})
            continue
        except (ValidationError, sa.exc.SQLAlchemyError) as exc:  # a row the DB refuses must not lose the file
            job.status = ImportJobStatus.FAILED
            job.error_report = [*job.error_report, {"row": row.get("_row"), "field": "_", "message": str(exc)[:400], "kind": "error"}]
            await session.flush()
            raise ValidationFailed(f"تعذّر إنشاء الصف رقم {row.get('_row')}", details={"row": row.get("_row")}) from exc
        created += 1
        _ = lead

    job.status = ImportJobStatus.COMPLETED
    job.committed_by = actor.id
    job.committed_at = datetime.now(UTC)
    job.stats = {
        **job.stats,
        "created": created,
        "skipped_duplicate": int(job.stats.get("skipped_duplicate", 0)) + still_duplicate,
        "errors": int(job.stats.get("errors", 0)) + len([f for f in failures if f["kind"] == "error"]),
    }
    if failures:
        job.error_report = [*job.error_report, *failures]
    await session.flush()
    await audit.write(
        session,
        actor=actor,
        action="lead_import.committed",
        entity_type="lead_import_job",
        entity_id=job.id,
        after={"filename": job.filename, "created": created, "skipped": job.stats["skipped_duplicate"], "rows": job.total_rows},
        ip=ip,
    )
    return job
