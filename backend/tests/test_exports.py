"""`POST /exports/{kind}` (spec §7) — permission, scope, workbook contents, flag and audit (§6).

The invariant these tests protect: a downloaded file never contains a row the caller could not list on
screen, and every download leaves an audit trail. Reading the workbook back with openpyxl (instead of
asserting on bytes) is what makes the column order and the Arabic headers part of the contract.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from tests.helpers import login

ADMIN = "admin@brokerhub.test"
MANAGER = "manager@brokerhub.test"
REP_A = "rep_a@brokerhub.test"

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(payload: bytes) -> list[list[object]]:
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook.worksheets[0]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


async def test_export_requires_reports_export(client) -> None:
    """A sales rep has no export permission at all: 403, not a truncated file (§6)."""
    async with login(client, REP_A) as auth:
        refused = await client.post("/api/v1/exports/requests", headers=auth, json={})
    assert refused.status_code == 403
    assert refused.json()["code"] in {"forbidden", "permission_denied"}


async def test_requests_export_matches_the_callers_own_list(client) -> None:
    """The file's row count equals what the same user lists, so scoping cannot drift per surface."""
    async with login(client, MANAGER) as auth:
        listed = (await client.get("/api/v1/requests?limit=1", headers=auth)).json()
        response = await client.post("/api/v1/exports/requests", headers=auth, json={})

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["x-export-truncated"] == "0"
    rows = _sheet(response.content)
    assert rows[0][:4] == ["رقم الطلب", "الحالة", "الأولوية", "القناة"]
    assert len(rows) - 1 == listed["total"]
    assert response.headers["x-export-rows"] == str(listed["total"])
    # the request number is the first column, and it is an integer in the sheet, not a string
    if len(rows) > 1:
        assert isinstance(rows[1][0], int)


async def test_clients_export_is_scoped_like_the_clients_list(client) -> None:
    async with login(client, MANAGER) as auth:
        manager_total = (await client.get("/api/v1/clients?limit=1", headers=auth)).json()["total"]
        workbook = await client.post("/api/v1/exports/clients", headers=auth, json={})
    rows = _sheet(workbook.content)
    assert rows[0][:3] == ["الشركة", "النوع", "الحالة"]
    assert len(rows) - 1 == manager_total


async def test_client_filter_narrows_the_file(client, ids: dict[str, dict[str, str]]) -> None:
    """`client_id` is honoured: the sheet only ever names that company."""
    one_client = ids["clients"]["ألف"]
    async with login(client, MANAGER) as auth:
        expected = (await client.get(f"/api/v1/clients/{one_client}", headers=auth)).json()["name"]
        response = await client.post("/api/v1/exports/requests", headers=auth, json={"client_id": one_client})
    rows = _sheet(response.content)
    names = {row[5] for row in rows[1:]}
    assert names in (set(), {expected})


async def test_row_filters_are_honoured(client) -> None:
    async with login(client, MANAGER) as auth:
        overdue = await client.post("/api/v1/exports/requests", headers=auth, json={"only_overdue": True})
        everything = await client.post("/api/v1/exports/requests", headers=auth, json={})
    overdue_rows = len(_sheet(overdue.content)) - 1
    all_rows = len(_sheet(everything.content)) - 1
    assert overdue_rows <= all_rows


async def test_report_export_is_flag_gated_and_equals_the_screen(client) -> None:
    """`insurer-performance` is Phase 2: 503 while the flag is off, then the same rows as the screen."""
    from app.core.db import get_sessionmaker
    from app.models.masters import FeatureFlag

    maker = get_sessionmaker()
    async with maker() as session:
        session.add(FeatureFlag(key="phase2.enabled", enabled=False, description="Phase 2 screens"))
        await session.commit()

    async with login(client, MANAGER) as auth, login(client, ADMIN) as admin:
        blocked = await client.post("/api/v1/exports/insurer-performance", headers=auth, json={})
        assert blocked.status_code == 503
        assert blocked.json()["code"] == "feature_disabled"

        await client.patch("/api/v1/feature-flags/phase2.enabled", json={"enabled": True}, headers=admin)
        try:
            exported = await client.post("/api/v1/exports/insurer-performance", headers=auth, json={})
            screen = (await client.get("/api/v1/reports/insurer-performance", headers=auth)).json()
        finally:
            await client.patch("/api/v1/feature-flags/phase2.enabled", json={"enabled": False}, headers=admin)

    assert exported.status_code == 200, exported.text
    rows = _sheet(exported.content)
    assert rows[0][:3] == ["شركة التأمين", "الطلبات", "متوسط أول رد (دقيقة)"]
    assert len(rows) - 1 == len(screen)
    if rows[1:]:
        assert rows[1][0] == screen[0]["insurer"]
        assert rows[1][4] in {"ممتاز", "مقبول", "يحتاج تدخلاً"}


async def test_every_download_is_audited(client) -> None:
    """The list endpoint stays privacy-safe, so the payload itself is read from the row."""
    from app.core.db import get_sessionmaker
    from app.models.requests import AuditLog
    from sqlalchemy import select

    async with login(client, MANAGER) as auth:
        await client.post("/api/v1/exports/requests", headers=auth, json={})
        mine = (await client.get("/api/v1/auth/audit/mine", headers=auth)).json()

    entries = mine if isinstance(mine, list) else mine["items"]
    assert "export.requests" in [row["action"] for row in entries]

    async with get_sessionmaker()() as session:
        row = (
            (await session.execute(select(AuditLog).where(AuditLog.action == "export.requests").order_by(AuditLog.occurred_at.desc())))
            .scalars()
            .first()
        )
    assert row is not None
    assert row.entity_type == "export"
    assert row.after is not None and row.after["rows"] >= 0
    assert "sheet" in row.after and "filters" in row.after


async def test_unknown_kind_is_rejected_before_any_query(client) -> None:
    async with login(client, MANAGER) as auth:
        response = await client.post("/api/v1/exports/salaries", headers=auth, json={})
    assert response.status_code == 422
